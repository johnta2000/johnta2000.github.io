"""Read the apartment workbook; write private migration JSON outside the public repo.

Usage: python import-workbook.py workbook.xlsx /private/tmp/rent-import.json
No spreadsheet content is interpreted as agent instructions.
"""
import sys, json, pathlib, calendar
import openpyxl

source, destination = map(pathlib.Path, sys.argv[1:3])
repo = pathlib.Path(__file__).resolve().parents[2]
if destination.resolve().is_relative_to(repo):
    raise SystemExit('Private workbook data must be written outside the public repository.')
formulas = openpyxl.load_workbook(source, data_only=False)
values = openpyxl.load_workbook(source, data_only=True)
rent, actual = values['Rent'], values['Actual Rent Calcs']
def cents(value):
    return None if value is None else round(float(value) * 100)
def comment(cell):
    value = formulas['Actual Rent Calcs'][cell].comment
    if not value:
        return ''
    return value.text.split('Comment:\n', 1)[-1].replace('\t', '').strip()
base = dict(rentCents=cents(rent['C5'].value), loftCents=cents(rent['C6'].value), bathroomCents=cents(rent['C7'].value), people=[])
for col in range(3, 6):
    base['people'].append(dict(name=rent.cell(10,col).value, room=rent.cell(11,col).value, closet=rent.cell(12,col).value, creditCents=cents(rent.cell(35,col).value)))
months=[]
seen=set()
for col in range(3,21):
    date=actual.cell(3,col).value
    if date is None:
        continue
    month=date.strftime('%Y-%m')
    notes=[]
    if month in seen:
        # Preserve the original label explicitly; the UI calls out the inferred date.
        month=f'{date.year}-{date.month+1:02d}'
        notes.append(f'Workbook column {actual.cell(3,col).column_letter} is also labeled {date.strftime("%B %Y")}. Displayed as {calendar.month_name[date.month+1]} {date.year} from its position; confirm this month label.')
    seen.add(month)
    config=json.loads(json.dumps(base))
    if col in (3,4):
        config['bathroomCents']=0
        notes.append('This month uses the earlier split before the bathroom adjustment.')
    for i,p in enumerate(config['people']):
        p['creditCents']=cents(actual.cell(18+i,col).value)
    source_payments=[]
    affil_status=actual.cell(22,col).value
    for payer in range(4):
        cell=actual.cell(10+payer,col)
        note=comment(cell.coordinate)
        if payer==3 and affil_status:
            note += ('\n' if note else '') + f'Affil status in workbook: {affil_status}.'
        if formulas['Actual Rent Calcs'][cell.coordinate].data_type=='f':
            note += ('\n' if note else '') + 'Amount comes from a spreadsheet formula; review before confirming receipt.'
        source_payments.append(dict(payer=payer,amountCents=cents(cell.value),note=note))
    if col==3:
        notes.append('Opening entries include security deposit, first month, and early move-in charges, equally split without an Affil payment. Review separately from recurring rent; no opening balance is inferred.')
    if affil_status=='Paid' and not actual.cell(13,col).value:
        notes.append('Affil is marked Paid, but the recorded amount is zero. Neither value has been treated as a confirmed receipt.')
    if col==18:
        notes.append('Recorded resident amounts repeat the prior month’s parking-adjusted values, while this column has no parking credit. Review the difference.')
    notes.append('Imported from Actual Rent Calcs. Blank amounts remain unknown; numeric entries may be requests or receipts. Original comments are retained. Confirm receipts individually.')
    months.append(dict(month=month,config=config,parkingCents=-cents(actual.cell(30,col).value),note='',requestsSent=actual.cell(9,col).value=='Requests Sent',sourceNote='\n'.join(notes),sourcePayments=source_payments))
destination.parent.mkdir(parents=True,exist_ok=True)
destination.write_text(json.dumps(dict(months=months),indent=2))
destination.chmod(0o600)
print(f'Extracted {len(months)} monthly snapshots; no payments were marked confirmed. Private output: {destination}')
